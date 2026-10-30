from pathlib import Path
from openpyxl import load_workbook

base = Path(r'c:\Users\aj1057\Desktop\NC')
path = base / 'NC_Library System_Branches_Manual Check _9_20_2026(1).xlsx'
report = base / 'workbook_report.txt'

lines = []
lines.append(f'path_exists={path.exists()}')

if path.exists():
    wb = load_workbook(path, data_only=True, read_only=True)
    lines.append(f'sheetnames={wb.sheetnames}')
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f'sheet={sheet}; rows={ws.max_row}; cols={ws.max_column}')
        rows = ws.iter_rows(min_row=1, max_row=min(15, ws.max_row), values_only=True)
        for idx, row in enumerate(rows, start=1):
            lines.append(f'{sheet}:{idx}:{row}')

report.write_text('\n'.join(lines), encoding='utf-8')
print('REPORT_PATH=' + str(report))
print('\n'.join(lines[:20]))
