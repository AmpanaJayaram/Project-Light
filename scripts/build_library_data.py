"""Build data/library-data.json from the NC branches workbook.

Coordinates, phone numbers and weekly hours come from the IMLS-based dataset
embedded in nc-library-finder-standalone.html. Workbook rows are matched to
those records by ZIP + street address, then by ZIP + name; anything left over
is placed at its ZIP centroid.

Usage: .venv/Scripts/python.exe scripts/build_library_data.py
"""
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
WORKBOOK = BASE / 'NC_Library System_Branches_Manual Check _9_20_2026(1).xlsx'
STANDALONE = BASE / 'nc-library-finder-standalone.html'
OUT = BASE / 'data' / 'library-data.json'

# Workbook system URLs that no longer work (checked 2026-10-29).
URL_FIXES = {
    # LibGuides site returns 404; NC Cardinal catalog page for the main branch works.
    'https://bladenco.libguides.com/': 'https://nextsandhill.nccardinal.org/eg/opac/library/BLADEN_MAIN',
    # Old domain redirects here.
    'https://granville.lib.nc.us/': 'https://www.granvillecounty.org/157/Library-System',
}

STREET_ABBR = {
    'street': 'st', 'avenue': 'ave', 'road': 'rd', 'drive': 'dr', 'boulevard': 'blvd',
    'highway': 'hwy', 'lane': 'ln', 'court': 'ct', 'place': 'pl', 'parkway': 'pkwy',
    'north': 'n', 'south': 's', 'east': 'e', 'west': 'w', 'suite': 'ste', 'circle': 'cir',
}
NAME_STOPWORDS = {'library', 'public', 'branch', 'county', 'memorial', 'regional', 'the', 'of', 'and'}


def norm_addr(value):
    words = re.sub(r'[^a-z0-9 ]', ' ', str(value or '').lower()).split()
    return ' '.join(STREET_ABBR.get(w, w) for w in words)


def name_tokens(value):
    words = re.sub(r'[^a-z0-9 ]', ' ', str(value or '').lower()).split()
    return {w for w in words if w not in NAME_STOPWORDS}


def slug(value):
    return re.sub(r'[^a-z0-9]+', '-', value.lower()).strip('-')


def load_embedded_data():
    html = STANDALONE.read_text(encoding='utf-8')
    match = re.search(r'const DATA = (\{.*?\});?\s*$', html, re.M)
    return json.loads(match.group(1))


def system_resources(stats):
    """Resource tags backed by the system's IMLS counts."""
    if not stats:
        return ['books']
    tags = []
    if stats.get('books'):
        tags.append('books')
    if stats.get('computers'):
        tags.append('computer')
    # IMLS leaves Wi-Fi blank (never 0) for systems that did not report it, so a
    # blank means unknown; keep the tag and flag it via unverified_resources().
    if stats.get('wifi') or stats.get('wifi') is None:
        tags.append('wifi')
    if stats.get('dbs'):
        tags.append('research')
    if stats.get('kids'):
        tags.extend(['kids', 'family'])
    if stats.get('teen'):
        tags.append('teens')
    if stats.get('adult'):
        tags.append('adults')
    return tags or ['books']


def unverified_resources(stats):
    if stats and stats.get('wifi') is None:
        return ['wifi']
    return []


def system_highlights(stats):
    if not stats:
        return ['Library services']
    out = []
    if stats.get('programs'):
        out.append(f"{stats['programs']:,} programs a year")
    if stats.get('computers'):
        out.append(f"{stats['computers']} public computers")
    if stats.get('ebooks'):
        out.append(f"{stats['ebooks']:,} e-books")
    return out or ['Library services']


def main():
    data = load_embedded_data()
    libs, systems, zips = data['libs'], data['sys'], data['zips']

    by_zip = defaultdict(list)
    for lib in libs:
        by_zip[lib['z']].append(lib)

    ws = load_workbook(WORKBOOK, data_only=True, read_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    # First pass: match rows to IMLS records so each workbook system learns its IMLS id.
    matched = []
    for row in rows:
        if not row or not row[3]:
            continue
        sys_id, sys_name, sys_url, out_id, out_name, out_type, county, address, city, zipcode = row
        zip5 = str(zipcode or '').strip()[:5].zfill(5) if zipcode else ''
        candidates = by_zip.get(zip5, [])
        addr = norm_addr(address)
        hit = next((c for c in candidates if norm_addr(c['a']) == addr), None)
        if not hit:
            tokens = name_tokens(out_name)
            scored = [(len(tokens & name_tokens(c['n'])), c) for c in candidates]
            scored = [s for s in scored if s[0] > 0]
            if scored:
                hit = max(scored, key=lambda s: s[0])[1]
        matched.append((row, zip5, hit))

    imls_system_for = {}
    for row, _, hit in matched:
        if hit:
            imls_system_for.setdefault(row[0], hit['s'])

    records, unplaced = [], []
    counts = {'address_or_name': 0, 'zip_centroid': 0}
    for row, zip5, hit in matched:
        sys_id, sys_name, sys_url, out_id, out_name, out_type, county, address, city, zipcode = row
        stats = systems.get(imls_system_for.get(sys_id, ''))
        if hit:
            lat, lng = hit['la'], hit['lo']
            counts['address_or_name'] += 1
        elif zip5 in zips:
            lat, lng = zips[zip5]
            counts['zip_centroid'] += 1
        else:
            unplaced.append(out_name)
            continue

        # System-wide program counts only describe full-service branches.
        if out_type.startswith('specialized'):
            resources = ['books', 'research']
        elif out_type == 'mobile' or out_type.startswith('self_service') or out_type == 'automated_kiosk':
            resources = ['books']
        else:
            resources = system_resources(stats)
        full_service = resources == system_resources(stats)
        records.append({
            'id': slug(out_id),
            'outletId': out_id,
            'name': out_name,
            'system': sys_name,
            'systemId': sys_id,
            'url': URL_FIXES.get(sys_url, sys_url),
            'type': out_type,
            'address': address,
            'city': city,
            'zip': zip5,
            'county': county,
            'lat': round(lat, 5),
            'lng': round(lng, 5),
            'phone': hit.get('p') if hit and hit.get('p') else 'See library website',
            'resources': resources,
            'unverifiedResources': unverified_resources(stats) if full_service else [],
            'programHighlights': system_highlights(stats),
            'hours': f"{hit['h']} hrs/week" if hit and hit.get('h') else 'Hours vary',
            'locationPrecision': 'exact' if hit else 'zip',
        })

    OUT.parent.mkdir(exist_ok=True)
    payload = {
        'source': WORKBOOK.name,
        'libraries': records,
        'zips': zips,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
    print(f'Wrote {len(records)} libraries to {OUT}')
    counties_out = OUT.parent / 'nc-counties.geojson'
    counties_out.write_text(json.dumps(data['counties']), encoding='utf-8')
    print(f"Wrote {len(data['counties']['features'])} county outlines to {counties_out}")
    print(f'Placement: {counts}')
    if unplaced:
        print(f'Skipped (no coordinates): {unplaced}')


if __name__ == '__main__':
    main()
