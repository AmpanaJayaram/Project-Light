from pathlib import Path
from openpyxl import load_workbook

base = Path(r'c:\Users\aj1057\Desktop\NC')
workbook_path = base / 'NC_Library System_Branches_Manual Check _9_20_2026(1).xlsx'
report_path = base / 'workbook_report.txt'

lines = []
lines.append(f'path_exists={workbook_path.exists()}')

if workbook_path.exists():
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    lines.append(f'sheetnames={wb.sheetnames}')
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f'sheet={sheet_name}; rows={ws.max_row}; cols={ws.max_column}')
        for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(10, ws.max_row), values_only=True), start=1):
            lines.append(f'{sheet_name}:{row_index}:{row}')

report_path.write_text('\n'.join(lines), encoding='utf-8')
print(f'WROTE {report_path}')
