from openpyxl import load_workbook
from pathlib import Path

path = Path(r'c:\Users\aj1057\Desktop\NC\NC_Library System_Branches_Manual Check _9_20_2026(1).xlsx')
print('exists:', path.exists())
wb = load_workbook(path, data_only=True, read_only=True)
print('sheets:', wb.sheetnames)
ws = wb.active
print('max_row', ws.max_row, 'max_col', ws.max_column)
for i in range(1, 11):
    row = [cell.value for cell in ws[i]]
    print(row)
