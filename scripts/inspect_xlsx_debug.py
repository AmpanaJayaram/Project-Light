from openpyxl import load_workbook
from pathlib import Path

base_dir = Path(__file__).resolve().parent.parent
workbook_path = base_dir / 'NC_Library System_Branches_Manual Check _9_20_2026(1).xlsx'
out_path = base_dir / 'workbook_inspection.txt'

lines = []
lines.append(f'exists: {workbook_path.exists()}')

if workbook_path.exists():
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    lines.append(f'sheets: {wb.sheetnames}')
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f'--- {sheet_name} rows={ws.max_row} cols={ws.max_column}')
        for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=True), start=1):
            lines.append(f'{idx}: {row}')

out_path.write_text('\n'.join(lines), encoding='utf-8')
print(f'Wrote inspection report to {out_path}')
